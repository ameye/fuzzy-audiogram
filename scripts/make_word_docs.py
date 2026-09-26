"""Build the Word versions of the submission documents.

Elsevier accepts DOCX, and tables and figure captions are wanted as separate files.
Everything here is built with python-docx rather than converted from the PDFs, so the
output carries real styles and editable tables.

House style throughout: Bookman Old Style 12 pt, justified, no first-line indent.

Run:  python3 scripts/make_word_docs.py
Writes: submission/*.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches

ROOT = Path(__file__).resolve().parent.parent
MS = ROOT / "manuscript"
SUB = ROOT / "submission"
QMD = MS / "manuscript_v4.qmd"

TITLE = ("A Mamdani fuzzy inference framework for graded pure-tone audiometric "
         "classification: development, validation, and open-source implementation")
AUTHOR = "Sanyaolu Ameye"
BYLINE = "Sanyaolu Ameye – MBBS, FWACS, FMCORL, Pg Cert AI/ML, MDS"
ROLE = "Consultant Otorhinolaryngology Head and Neck Surgeon"
ORCID = "0000-0002-5217-7997"
EMAIL = "sanyaameye@hotmail.com"
REPO = "https://github.com/ameye/fuzzy-audiogram"
AFFIL = ROLE
FONT = "Bookman Old Style"

HIGHLIGHTS = [
    "Mamdani fuzzy inference provides graded classification of pure-tone audiograms.",
    "Ruspini partition ensures memberships sum to unity across 0 to 120 dB HL.",
    "Validated on 19,623 NHANES adult ears with participant-level splitting.",
    "Reference label shown to be a deterministic thresholding of its input.",
    "Open-source implementation with a REST API for EHR integration.",
]


def new_doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    st.paragraph_format.first_line_indent = Inches(0)
    st.paragraph_format.space_after = Pt(8)
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)
    # make sure the theme font does not override Bookman in Word
    try:
        d.styles["Normal"].element.rPr.rFonts.set(
            __import__("docx.oxml.ns", fromlist=["qn"]).qn("w:eastAsia"), FONT)
    except Exception:
        pass
    return d


def para(d, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         space_after=8, italic=False):
    p = d.add_paragraph()
    p.paragraph_format.alignment = align
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    return p


def heading(d, text, size=13):
    return para(d, text, size=size, bold=True,
                align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)


def md_to_runs(p, text):
    """add text to a paragraph, honouring **bold** and *italic*"""
    for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            r = p.add_run(chunk[2:-2]); r.bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            r = p.add_run(chunk[1:-1]); r.italic = True
        else:
            r = p.add_run(chunk)
        r.font.name = FONT
        r.font.size = Pt(12)


def add_table(d, rows):
    hdr = rows[0]
    t = d.add_table(rows=len(rows), cols=len(hdr))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            if j >= len(hdr):
                continue
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(cell)
            r.font.name = FONT
            r.font.size = Pt(10)
            r.bold = (i == 0)
    return t


# ---------------------------------------------------------------- documents

def cover_letter(path):
    d = new_doc()
    heading(d, "Cover letter", 15)
    para(d, "To:", align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)
    para(d, "The Editors-in-Chief", align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)
    para(d, "Computer Methods and Programs in Biomedicine",
         align=WD_ALIGN_PARAGRAPH.LEFT, space_after=12)
    para(d, f"Manuscript title: {TITLE}", align=WD_ALIGN_PARAGRAPH.LEFT, space_after=12)

    para(d, "Dear Editors,")
    para(d, "I am submitting for consideration a research article describing an open-source "
            "Mamdani fuzzy inference system for grading pure-tone audiograms. The work "
            "replaces the fixed severity boundaries that audiometry currently relies on with a "
            "graded index built on a formal Ruspini partition, and validates it on 19,623 adult "
            "ears drawn from three NHANES cycles with a participant-level train and test split.")
    para(d, "The contribution I would ask you to weigh is not a new classifier but a "
            "demonstrable and reproducible construction, which is why I have submitted to CMPB "
            "rather than to a clinical audiology journal. Three features make the method "
            "auditable rather than merely reported. The six trapezoidal membership functions "
            "are built by a published algorithm in which adjacent cores never move and the "
            "shoulders are proven to sum to one across the whole universe, so the partition "
            "property is a property of the construction and not a numerical coincidence. Every "
            "threshold, table value and figure in the manuscript is produced by committed code, "
            "and the repository holds the verification scripts alongside the pipeline. And the "
            "analysis reports where the method does not win: the reference grade is a "
            "deterministic thresholding of the pure-tone average, which a proportional-odds "
            "model recovers exactly, and a seven-frequency gradient-boosting model outperforms "
            "the index on every metric. The manuscript states both and locates the index's value "
            "accordingly, in interpretability, in a graded output no crisp classifier provides, "
            "and in explicit behaviour at the boundaries where graders disagree.")
    para(d, "The evaluation is deliberately adversarial towards the method. Agreement is "
            "stratified by distance from the severity boundary, where the index should help and "
            "where a summary statistic hides the effect. The transition width is swept rather "
            "than reported at a single setting. A second classification schema is run end to end "
            "to establish that the results depend on the method rather than on the choice of "
            "reference standard. And the borderline and clear-case agreement rates are reported "
            "with confidence intervals, so the reader can see the range in which the index is "
            "and is not informative.")
    para(d, "The system, the analysis pipeline and the verification scripts are available as an "
            f"open-source implementation at {REPO}, and a web calculator with a documented REST "
            "API for EHR integration is included for clinical evaluation.")
    para(d, "This manuscript is original, has not been published previously, and is not under "
            "consideration elsewhere. The author declares no competing interests and received no "
            "specific funding for this work. A declaration of generative AI use is included in "
            "the manuscript.")
    para(d, "Thank you for considering this submission.")
    para(d, "Yours sincerely,", space_after=16)
    para(d, BYLINE, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)
    para(d, ROLE, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)
    para(d, f"ORCID: {ORCID}", align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)
    para(d, EMAIL, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=0)
    d.save(path)


def title_page(path):
    d = new_doc()
    para(d, TITLE, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=14)
    para(d, BYLINE, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    para(d, ROLE, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    para(d, f"ORCID: {ORCID}", align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    para(d, f"Corresponding author: {AUTHOR}, {EMAIL}",
         align=WD_ALIGN_PARAGRAPH.LEFT, space_after=12)

    heading(d, "Article type"); para(d, "Full-length research article")
    heading(d, "Keywords")
    para(d, "audiometry, hearing loss classification, fuzzy inference, Ruspini partition, "
            "reproducibility, clinical decision support")
    heading(d, "Highlights")
    for h in HIGHLIGHTS:
        para(d, h, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
    heading(d, "Author contributions (CRediT)")
    para(d, f"{AUTHOR}: Conceptualization, Methodology, Software, Validation, Formal analysis, "
            "Investigation, Data curation, Writing - original draft, Writing - review and "
            "editing, Visualization.")
    heading(d, "Funding")
    para(d, "This research received no specific grant from any funding agency in the public, "
            "commercial or not-for-profit sectors.")
    heading(d, "Declaration of competing interest")
    para(d, "The author reports no competing interests to declare.")
    heading(d, "Data availability")
    para(d, "The NHANES data analysed here are publicly available from the U.S. Centers for "
            "Disease Control and Prevention. The fuzzy inference system, the Ruspini membership "
            f"function construction, the analysis pipeline and the verification scripts are "
            f"available at {REPO}.")
    heading(d, "Declaration of generative AI use")
    para(d, "During the preparation of this work the author used a large language model "
            "assistant to support drafting, language editing and code documentation. No AI or "
            "machine-learning tool generated the underlying data, selected the cohort, or "
            "computed any reported statistic; the fuzzy inference system is a transparent, "
            "rule-based model with no learned parameters. The author reviewed and verified all "
            "content and takes full responsibility for the work.")
    heading(d, "Counts")
    para(d, "Main text 3,497 words (CMPB limit 3,500 excluding abstract, references and figure "
            "captions). Abstract 277 words (limit 350). 18 references. 4 figures, 4 tables, "
            "1 supplementary file.")
    d.save(path)


def highlights(path):
    d = new_doc()
    heading(d, "Highlights", 15)
    para(d, TITLE, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    para(d, AUTHOR, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    para(d, "Each highlight is under 85 characters including spaces, as required.",
         size=10, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=12)
    for h in HIGHLIGHTS:
        para(d, h, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)
    d.save(path)


def parse_md_tables(text):
    """pull the markdown pipe tables out, in order"""
    tables, cur = [], []
    for line in text.splitlines():
        if line.strip().startswith("|"):
            cur.append(line.strip())
        else:
            if cur:
                tables.append(cur); cur = []
    if cur:
        tables.append(cur)
    out = []
    for t in tables:
        rows = []
        for r in t:
            cells = [c.strip() for c in r.strip("|").split("|")]
            if all(set(c) <= set("-: ") for c in cells):
                continue
            rows.append(cells)
        if rows:
            out.append(rows)
    return out


def tables_doc(path):
    text = QMD.read_text()
    body = text[text.find("# Tables"):text.find("# Figure Captions")]
    # captions and their tables, in document order
    caps = re.findall(r"\*\*Table (\d)\.\*\*\s*(.*?)(?=\n\n|\n\*\*)", body, re.S)
    tbls = parse_md_tables(body)
    d = new_doc()
    heading(d, "Tables", 15)
    para(d, TITLE, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    para(d, AUTHOR, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=14)
    made = 0
    for i, rows in enumerate(tbls):
        n = i + 1
        cap = ""
        for cn, ct in caps:
            if int(cn) == n:
                cap = ct.strip()
                break
        p = d.add_paragraph(); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Inches(0); p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"Table {n}. "); r.bold = True; r.font.name = FONT; r.font.size = Pt(12)
        md_to_runs(p, cap)
        add_table(d, rows)
        para(d, "", space_after=10)
        made += 1
    d.save(path)
    return made


def captions_doc(path):
    text = QMD.read_text()
    body = text[text.find("# Figure Captions"):]
    caps = re.findall(r"\*\*Figure (\d)\.\*\*\s*(.*?)(?=\n\s*\n|\n\*\*Figure|$)", body, re.S)
    d = new_doc()
    heading(d, "Figure captions", 15)
    para(d, TITLE, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    para(d, AUTHOR, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=14)
    for n, cap in caps:
        p = d.add_paragraph(); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(0); p.paragraph_format.space_after = Pt(12)
        r = p.add_run(f"Figure {n}. "); r.bold = True; r.font.name = FONT; r.font.size = Pt(12)
        md_to_runs(p, " ".join(cap.split()))
    d.save(path)
    return len(caps)



def supplementary_doc(path):
    """Supplementary material as Word: Table S1 and Figure S1."""
    import json
    from docx.enum.section import WD_ORIENT

    ORDER = ["normal", "mild", "moderate", "moderately_severe", "severe", "profound"]
    LABEL = {"normal": "Normal", "mild": "Mild", "moderate": "Moderate",
             "moderately_severe": "Moderately severe", "severe": "Severe",
             "profound": "Profound"}
    CALIB = [23.8, 48.9, 63.8, 66.4, 95.0]
    p = json.loads((ROOT / "data" / "output_participant"
                    / "params_participant.json").read_text())

    d = new_doc()
    # landscape: Table S1 has eight columns
    for s in d.sections:
        w, h = s.page_width, s.page_height
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = h, w
        s.left_margin = s.right_margin = Inches(0.8)

    heading(d, "Supplementary material", 15)
    para(d, TITLE, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    para(d, f"{AUTHOR} | ORCID: {ORCID}", align=WD_ALIGN_PARAGRAPH.LEFT, space_after=10)
    para(d, "Ruspini partition parameters and the temporal tracking demonstration. "
            "Generated by scripts/make_supplementary.py from the fitted parameters and the "
            "live temporal module.", size=10, italic=True, space_after=12)

    heading(d, "Table S1. Ruspini partition parameters")
    hdr = ["Category", "P25 core (dB HL)", "P75 core (dB HL)", "Trapezoid [a, b, c, d]",
           "Transition band (dB HL)", "Width (dB)", "0.5 crossover (dB HL)",
           "Calibrated cut point (dB HL)"]
    rows = [hdr]
    for i, k in enumerate(ORDER):
        a, b, c, dd = p[k]
        if i == 0:
            band = width = cross = cut = "\u2014"
        else:
            pa = p[ORDER[i - 1]]
            band = f"{pa[2]:.1f}\u2013{pa[3]:.1f}"
            width = f"{pa[3] - pa[2]:.1f}"
            cross = f"{(pa[2] + b) / 2:.1f}"
            cut = f"{CALIB[i - 1]:.1f}"
        rows.append([LABEL[k], f"{b:.1f}", f"{c:.1f}", f"[{a:.1f}, {b:.1f}, {c:.1f}, {dd:.1f}]",
                     band, width, cross, cut])
    add_table(d, rows)
    para(d, "The trapezoid [a, b, c, d] is the membership function for that category, with the "
            "core at [b, c] held at the empirical P25 and P75 and shoulders set so adjacent "
            "functions are complementary. Transition bands and 0.5 crossovers are geometric "
            "consequences of the parameters; the calibrated label cut points are fitted "
            "separately and are not the crossovers.", size=10, space_after=14)

    heading(d, "Figure S1. Temporal tracking under ototoxic drift")
    img = MS / "supplementary" / "Figure_S1.png"
    if img.exists():
        pp = d.add_paragraph()
        pp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.paragraph_format.space_after = Pt(8)
        pp.add_run().add_picture(str(img), width=Inches(9.2))
    para(d, "Serial audiograms (a) and the resulting FAI trajectory (b). Circles mark timepoints "
            "where the fuzzy-adapted ASHA shift criteria flag a significant change. Illustrative "
            "only: the cohort is repeated cross-sectional, so no longitudinal audiometric "
            "sub-cohort exists in these data.", size=10)
    d.save(path)


def main():
    SUB.mkdir(parents=True, exist_ok=True)
    cover_letter(SUB / "01_Cover_Letter.docx"); print("  01_Cover_Letter.docx")
    title_page(SUB / "02_Title_Page.docx"); print("  02_Title_Page.docx")
    highlights(SUB / "04_Highlights.docx"); print("  04_Highlights.docx")
    n = tables_doc(SUB / "07_Tables.docx"); print(f"  07_Tables.docx ({n} tables)")
    n = captions_doc(SUB / "08_Figure_Captions.docx"); print(f"  08_Figure_Captions.docx ({n} captions)")
    supplementary_doc(SUB / "05_Supplementary_Material.docx"); print("  05_Supplementary_Material.docx")


if __name__ == "__main__":
    main()
